"""Sinh báo cáo Word (bản gọn cho buổi trình bày 10 phút): python build_report.py"""
import json, os
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx_helpers import Report

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figures")
RES = os.path.join(HERE, "..", "demo", "results")
ev = json.load(open(os.path.join(RES, "eval.json"), encoding="utf8"))
cfg = json.load(open(os.path.join(RES, "config.json"), encoding="utf8"))

R = Report()
p, h, eq, fig, table, bullets = R.p, R.h, R.eq, R.fig, R.table, R.bullets
f4 = lambda v: f"{v:.3f}".replace(".", ",")

# =============================================================== TRANG BÌA
for txt, sz, b in [("[TÊN TRƯỜNG ĐẠI HỌC]", 14, True), ("[TÊN KHOA]", 13, True)]:
    par = R.doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run(txt); r.font.size = Pt(sz); r.bold = b
for _ in range(4):
    R.doc.add_paragraph()
par = R.doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = par.add_run("BÁO CÁO GIỮA KỲ – HỌC PHẦN: [TÊN HỌC PHẦN]"); r.font.size = Pt(14); r.bold = True
R.doc.add_paragraph()
par = R.doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = par.add_run("MÔ HÌNH HOÁ CHUỖI THỜI GIAN NHIỀU CHIỀU BẰNG HỖN HỢP GAUSSIAN ĐỘNG KẾT HỢP MẠNG NƠ-RON HỒI QUY (DGM²)")
r.font.size = Pt(18); r.bold = True
par = R.doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = par.add_run("Ứng dụng: dự báo khí hậu trên dữ liệu thưa USHCN"); r.font.size = Pt(14); r.italic = True
for _ in range(6):
    R.doc.add_paragraph()
for line in ["Giảng viên hướng dẫn: [Họ tên giảng viên]", "Sinh viên thực hiện: [Họ tên] – MSSV: [MSSV]", "Lớp: [Lớp]"]:
    par = R.doc.add_paragraph(); par.paragraph_format.left_indent = Cm(4)
    par.add_run(line).font.size = Pt(13)
for _ in range(6):
    R.doc.add_paragraph()
par = R.doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
par.add_run("[Địa điểm], tháng 9/2026").italic = True
R.page_break()

# =============================================================== MỤC LỤC
par = R.doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = par.add_run("MỤC LỤC"); r.bold = True; r.font.size = Pt(16)
R.toc()
p("(Nếu mục lục chưa hiển thị: trong Word nhấn chuột phải vào vùng này → Update Field → Update entire table.)",
  align="center", size=10)
R.page_break()

# =============================================================== CHƯƠNG 1
h("Chương 1. Giới thiệu")
h("1.1. Bài toán", 2)
p("Chuỗi thời gian nhiều chiều (multivariate time series – MTS) ghi lại đồng thời nhiều biến theo thời gian, ví dụ "
  "nhiệt độ, lượng mưa, độ dày tuyết của một trạm khí tượng, hay các chỉ số xét nghiệm của một bệnh nhân. Trong thực tế "
  "các biến được đo từ nhiều nguồn với tần suất khác nhau, nên khi ghép lên cùng một trục thời gian thì có rất nhiều ô "
  "**bị thiếu** – ta gọi là MTS **thưa**. Bài toán đặt ra: *dùng một đoạn quá khứ không đầy đủ để dự báo các giá trị "
  "tương lai*.")
p("Báo cáo tìm hiểu mô hình **DGM²** (*Dynamic Gaussian Mixture based Deep Generative Model*, Wu và cộng sự, AAAI 2021), "
  "một cách kết hợp **mô hình hỗn hợp Gaussian (GMM)** với **mạng nơ-ron hồi quy (LSTM)** để giải bài toán trên, và "
  "trình bày demo ứng dụng dự báo khí hậu trên bộ dữ liệu thưa USHCN của bài báo.")
h("1.2. Ý tưởng chính", 2)
p("Quan sát then chốt: các chuỗi thời gian khác nhau thường **chia sẻ những trạng thái ẩn giống nhau**. Hình "
  "[[fig:intro]] minh hoạ hai bệnh nhân lọc máu: mỗi thời điểm, bệnh nhân ở một trạng thái ẩn (khoẻ mạnh, rối loạn thận, "
  "thiếu máu), và vector đo được sinh ra từ trạng thái đó. Các vector cùng trạng thái – dù đến từ bệnh nhân khác nhau – "
  "nằm gần nhau, tạo thành **một cụm**. Với khí hậu cũng vậy: các ngày thuộc cùng một “kiểu thời tiết” có các giá trị "
  "nhiệt độ, mưa, tuyết tương tự nhau.")
fig(os.path.join(FIG, "fig1_latent_states.png"),
    "Cấu trúc cụm tiềm ẩn phía sau MTS thưa của hai bệnh nhân (Hình 1 trong bài báo). Ô vàng: giá trị quan sát được; "
    "ô trắng: giá trị bị thiếu.", 11, label="intro")
p("Từ đó DGM² không dự báo trực tiếp từng vector thưa $\\mathbf{x}_t \\to \\mathbf{x}_{t+1}$, mà dự báo **cụm** tiếp theo "
  "$z_t \\to z_{t+1}$, rồi sinh giá trị từ cụm đó. Mỗi cụm là một thành phần Gaussian được học từ **toàn bộ** dữ liệu, nên "
  "dù một chuỗi bị thiếu nhiều, mô hình vẫn dự báo tốt nhờ thông tin của các chuỗi khác. Ba ý tưởng cốt lõi:")
bullets([
    "**GMM** mô tả các cụm (trạng thái) trong không gian dữ liệu;",
    "**LSTM** học quy luật chuyển giữa các cụm theo thời gian;",
    "**Hỗn hợp Gaussian động**: trọng số của GMM thay đổi theo từng thời điểm theo dự đoán của LSTM.",
])

# =============================================================== CHƯƠNG 2
h("Chương 2. Cơ sở lý thuyết")
h("2.1. Biểu diễn chuỗi thời gian thưa", 2)
p("Một MTS gồm $w$ bước thời gian và $d$ biến được viết là $\\mathbf{x}_{1:w} = (\\mathbf{x}_1, \\dots, \\mathbf{x}_w)$, với "
  "$\\mathbf{x}_t = (x^1_t, \\dots, x^d_t)^\\top \\in \\mathbb{R}^d$ là vector các biến tại thời điểm $t$. Để đánh dấu giá trị nào "
  "quan sát được, dùng **mặt nạ** $m^i_t \\in \\{0, 1\\}$: $m^i_t = 1$ nếu $x^i_t$ đo được, $m^i_t = 0$ nếu bị thiếu. "
  "Mục tiêu là dự báo $r$ bước tương lai $\\tilde{\\mathbf{x}}_{w+1:w+r}$ khi biết $\\mathbf{x}_{1:w}$ và $\\mathbf{m}_{1:w}$. Trong "
  "demo, $d = 5$ biến khí hậu, $w = 80$ ngày và $r = 20$ ngày.")

h("2.2. Mô hình hỗn hợp Gaussian (GMM)", 2)
p("GMM giả định dữ liệu được sinh ra từ $k$ cụm, mỗi cụm là một phân phối Gaussian. DGM² dùng Gaussian **đẳng hướng** "
  "(cùng phương sai theo mọi chiều), khi đó mật độ của một vector $\\mathbf{x}$ là:")
eq(r"p(\mathbf{x}) = \sum_{c=1}^{k} \pi_c \, \mathcal{N}(\mathbf{x} \mid \boldsymbol{\mu}_c, \sigma^{-1}\mathbf{I})", label="gmm")
bullets([
    "$\\boldsymbol{\\mu}_c$: **tâm** của cụm $c$ – vector “điển hình” của cụm (ví dụ một kiểu thời tiết);",
    "$\\pi_c$: **trọng số trộn** – xác suất một điểm thuộc cụm $c$ ($\\pi_c \\ge 0$, $\\sum_c \\pi_c = 1$);",
    "$\\sigma^{-1}$: phương sai, cho biết các điểm trong cụm phân tán quanh tâm bao xa.",
])
p("Cách hiểu **sinh dữ liệu** của GMM gồm hai bước: (1) chọn một cụm $z \\in \\{1, \\dots, k\\}$ với xác suất $\\pi_z$; (2) sinh "
  "$\\mathbf{x}$ quanh tâm $\\boldsymbol{\\mu}_z$ theo phân phối Gaussian. Ngược lại, khi đã thấy $\\mathbf{x}$, xác suất nó thuộc "
  "cụm $c$ tính theo công thức Bayes:")
eq(r"p(z = c \mid \mathbf{x}) = \frac{\pi_c \, \mathcal{N}(\mathbf{x} \mid \boldsymbol{\mu}_c, \sigma^{-1}\mathbf{I})}{\sum_{c'=1}^{k} \pi_{c'} \, \mathcal{N}(\mathbf{x} \mid \boldsymbol{\mu}_{c'}, \sigma^{-1}\mathbf{I})}", label="resp")
p("Với Gaussian đẳng hướng, $\\log \\mathcal{N}(\\mathbf{x} \\mid \\boldsymbol{\\mu}_c, \\sigma^{-1}\\mathbf{I}) = -\\frac{\\sigma}{2}"
  "\\|\\mathbf{x} - \\boldsymbol{\\mu}_c\\|^2 + \\text{hằng số}$: điểm càng **gần tâm** thì càng có khả năng thuộc cụm đó.")
p("**Hạn chế khi áp dụng cho chuỗi thời gian:** trong GMM thường, trọng số $\\pi_c$ là **cố định**, không phụ thuộc thời "
  "gian. Nhưng với chuỗi thời gian, cụm của ngày mai phụ thuộc rõ vào các ngày trước (hôm nay đang rét thì ngày mai nhiều "
  "khả năng vẫn rét). DGM² giải quyết bằng cách cho trọng số trộn **thay đổi theo thời gian**, do một mạng LSTM dự đoán. "
  "LSTM là mạng nơ-ron hồi quy có trạng thái ẩn $\\mathbf{h}_t$ tóm tắt toàn bộ lịch sử đã đọc, nên nhớ được các phụ thuộc "
  "dài hạn.")

# =============================================================== CHƯƠNG 3
h("Chương 3. Mô hình DGM²")
h("3.1. Kiến trúc tổng quan", 2)
p("DGM² gồm ba khối được huấn luyện đồng thời (Hình [[fig:arch]]):")
bullets([
    "**Lớp tiền nội suy:** điền tạm các ô thiếu để mạng nơ-ron có đầu vào đầy đủ;",
    "**Mạng sinh** (Hình [[fig:arch]]b): LSTM dự đoán cụm tiếp theo, rồi sinh dữ liệu từ hỗn hợp Gaussian động – đây "
    "là phần dùng để **dự báo**;",
    "**Mạng suy diễn** (Hình [[fig:arch]]c): đọc dữ liệu quá khứ và suy ra mỗi thời điểm đang thuộc cụm nào – phần này "
    "dùng để **huấn luyện** và để “đọc” quá khứ trước khi dự báo.",
])
fig(os.path.join(FIG, "fig2_architecture.png"),
    "(a) Điều chỉnh động hỗn hợp Gaussian với hai thành phần; (b) mạng sinh; (c) mạng suy diễn, $\\gamma(\\cdot)$ là hàm "
    "cổng (Hình 2 trong bài báo).", 16, label="arch")

h("3.2. Lớp tiền nội suy", 2)
p("Một giá trị bị thiếu $x^i_{t^*}$ được ước lượng bằng **trung bình có trọng số** của các giá trị quan sát được của cùng "
  "biến ở những ngày lân cận; ngày càng gần $t^*$ thì trọng số càng lớn:")
eq(r"\bar{x}^i_{t^*} = \frac{\sum_{t=1}^{w} e^{-\alpha_i (t^* - t)^2} \, m^i_t \, x^i_t}{\sum_{t=1}^{w} e^{-\alpha_i (t^* - t)^2} \, m^i_t}", label="impute")
p("Thừa số $m^i_t$ loại bỏ các ngày bị thiếu; $\\alpha_i$ là tham số **học được** điều khiển độ rộng “cửa sổ” của từng "
  "biến (biến thay đổi chậm như nhiệt độ dùng cửa sổ rộng, biến biến động mạnh như lượng mưa dùng cửa sổ hẹp). Mẫu số "
  "gọi là **cường độ quan sát** – càng lớn thì càng có nhiều quan sát gần đó, giá trị điền càng đáng tin. Bài báo còn "
  "kết hợp thêm thông tin giữa các biến qua hệ số tương quan học được $\\rho_{ij}$. Đầu vào của mạng là: giữ nguyên giá "
  "trị quan sát được, chỉ thay các ô thiếu bằng $\\bar{x}$.")

h("3.3. Chuyển trạng thái cụm bằng LSTM", 2)
p("Mỗi thời điểm $t$ được gắn một **biến cụm** $z_t \\in \\{1, \\dots, k\\}$ cho biết $\\mathbf{x}_t$ thuộc cụm nào. Xác suất cụm "
  "của bước tiếp theo được dự đoán từ **toàn bộ lịch sử cụm** $z_{1:t}$ bằng một LSTM:")
eq(r"p(z_{t+1} \mid z_{1:t}) = \mathrm{softmax}\left(\mathrm{MLP}(\mathbf{h}_t)\right), \qquad \mathbf{h}_t = \mathrm{LSTM}(z_t, \mathbf{h}_{t-1})", label="trans")
p("LSTM nhận cụm hiện tại $z_t$, cập nhật trạng thái ẩn $\\mathbf{h}_t$ (bộ nhớ về các cụm đã qua); MLP và softmax biến "
  "$\\mathbf{h}_t$ thành một vector xác suất gồm $k$ phần tử có tổng bằng 1. Ví dụ, nếu 30 ngày qua chuỗi liên tục ở cụm "
  "“mùa đông lạnh” thì LSTM sẽ cho xác suất cao cho cụm này ở ngày tiếp theo.")

h("3.4. Hỗn hợp Gaussian động – thành phần cốt lõi", 2)
p("Đây là điểm mới quan trọng nhất của DGM². Thay cho trọng số cố định $\\pi_c$ của GMM thường, tại mỗi bước DGM² dùng "
  "trọng số **động**:")
eq(r"\boldsymbol{\psi}_{t+1} = (1 - \gamma) \, p(z_{t+1} \mid z_{1:t}) + \gamma \, p(\boldsymbol{\mu})", label="psi")
bullets([
    "$p(z_{t+1} \\mid z_{1:t})$ – **phần điều chỉnh động**: dự đoán của LSTM ở công thức ([[eq:trans]]), thay đổi theo từng "
    "chuỗi và từng thời điểm;",
    "$p(\\boldsymbol{\\mu}) = [p(\\boldsymbol{\\mu}_1), \\dots, p(\\boldsymbol{\\mu}_k)]$ – **hỗn hợp cơ sở**: tỉ lệ xuất hiện chung "
    "của mỗi cụm trên toàn bộ dữ liệu (ước lượng bằng trung bình xác suất cụm trên mỗi batch huấn luyện);",
    "$\\gamma \\in [0, 1]$ – hệ số pha trộn giữa hai phần.",
])
p("Khi đó dữ liệu ở bước $t+1$ được sinh ra từ một GMM có trọng số $\\boldsymbol{\\psi}_{t+1}$ – một GMM mà “hình dạng” thay "
  "đổi theo thời gian:")
eq(r"p(\mathbf{x}_{t+1} \mid z_{1:t}) = \sum_{c=1}^{k} \psi_{t+1}[c] \; \mathcal{N}(\mathbf{x}_{t+1} \mid \boldsymbol{\mu}_c, \sigma^{-1}\mathbf{I})", label="emis")
p("Hình [[fig:arch]](a) minh hoạ: hỗn hợp cơ sở có hai cụm (xanh, đỏ); tại bước $t+1$, LSTM dự đoán chuỗi thuộc cụm "
  "đỏ nên hỗn hợp bị “kéo” về phía cụm đỏ. **Vì sao cần giữ hỗn hợp cơ sở?** Nó gắn các cụm vào cấu trúc chung của "
  "toàn bộ dữ liệu, giúp các tâm $\\boldsymbol{\\mu}_c$ học ổn định và dự báo không “trôi” khi LSTM dự đoán sai. Hai trường "
  "hợp biên: $\\gamma = 1$ thì mô hình trở thành GMM tĩnh (mất động học); $\\gamma = 0$ thì chỉ còn dự đoán của LSTM. "
  "Bài báo cho thấy một giá trị nhỏ như $\\gamma = 0{,}01$ là tốt nhất.")
p("**Ví dụ số.** Giả sử $k = 3$, LSTM dự đoán $p(z_{t+1} \\mid z_{1:t}) = [0{,}7;\\ 0{,}2;\\ 0{,}1]$, hỗn hợp cơ sở "
  "$p(\\boldsymbol{\\mu}) = [0{,}3;\\ 0{,}3;\\ 0{,}4]$ và $\\gamma = 0{,}1$. Khi đó "
  "$\\psi_{t+1} = 0{,}9 \\cdot [0{,}7;\\ 0{,}2;\\ 0{,}1] + 0{,}1 \\cdot [0{,}3;\\ 0{,}3;\\ 0{,}4] = [0{,}66;\\ 0{,}21;\\ 0{,}13]$, và "
  "dự báo là $\\tilde{\\mathbf{x}}_{t+1} = 0{,}66\\,\\boldsymbol{\\mu}_1 + 0{,}21\\,\\boldsymbol{\\mu}_2 + 0{,}13\\,\\boldsymbol{\\mu}_3$ – chủ "
  "yếu theo cụm 1 nhưng vẫn “pha” một phần các cụm khác.", indent=False)

h("3.5. Mạng suy diễn", 2)
p("Để huấn luyện, cần biết mỗi thời điểm trong dữ liệu thật thuộc cụm nào – nhưng cụm là **ẩn**, không có nhãn. Mạng "
  "suy diễn ước lượng điều này bằng một LSTM thứ hai đọc dữ liệu (đã tiền nội suy):")
eq(r"q_\phi(z_t \mid \mathbf{x}_{1:t}, z_{t-1}) = \mathrm{softmax}\left(\mathrm{MLP}([\tilde{\mathbf{h}}_t, z_{t-1}])\right)", label="q")
eq(r"\tilde{\mathbf{h}}_t = \mathrm{LSTM}([\mathbf{x}_t, \mathbf{m}_t], \tilde{\mathbf{h}}_{t-1})", False)
p("Khác với mạng sinh (chỉ nhìn các cụm trước đó), mạng suy diễn **nhìn thấy cả dữ liệu** $\\mathbf{x}_{1:t}$, nên biết khá "
  "chính xác chuỗi đang ở cụm nào. Có thể hiểu nó đóng vai “giáo viên”: kết quả $q_\\phi$ được dùng để dạy mạng sinh dự "
  "đoán cụm, và để học các tâm cụm. Việc chọn cụm là rời rạc nên không lấy đạo hàm được; bài báo dùng thủ thuật "
  "**Gumbel-softmax** (lấy mẫu “mềm”) để cả mô hình huấn luyện được bằng lan truyền ngược.")

h("3.6. Hàm mục tiêu huấn luyện (ELBO)", 2)
p("Mục tiêu lý tưởng là cực đại xác suất sinh ra dữ liệu thật, nhưng việc này đòi hỏi cộng qua mọi chuỗi cụm có thể "
  "($k^w$ khả năng) nên không tính được. DGM² thay bằng cách cực đại một **cận dưới** (ELBO – evidence lower bound). "
  "Viết gọn, hàm mục tiêu gồm hai phần – phần (1) tái tạo và phần (2) khớp chuyển trạng thái:")
eq(r"\ell = \sum_{t} \sum_{c} \bar{q}_t[c] \, \log \mathcal{N}(\mathbf{x}_t \mid \boldsymbol{\mu}_c, \sigma^{-1}\mathbf{I})", label="elbo")
eq(r"- \sum_{t} \mathcal{D}_{KL}\left( q_\phi(z_t \mid \cdot) \,\|\, p(z_t \mid z_{1:t-1}) \right)", False)
p("trong đó $\\bar{q}_t = (1 - \\gamma)\\,q_\\phi(z_t \\mid \\cdot) + \\gamma\\,p(\\boldsymbol{\\mu})$ là trọng số cụm đã pha trộn như "
  "công thức ([[eq:psi]]). Ý nghĩa:", indent=False)
bullets([
    "**Phần (1) – tái tạo:** mỗi quan sát $\\mathbf{x}_t$ phải nằm gần tâm của cụm mà nó được gán. Phần này kéo các tâm "
    "$\\boldsymbol{\\mu}_c$ về giữa các nhóm dữ liệu – giống việc học một GMM. Chỉ tính trên các giá trị quan sát được "
    "(nhân với mặt nạ $m^i_t$), nên ô thiếu không ảnh hưởng.",
    "**Phần (2) – khớp chuyển trạng thái:** độ đo KL (khoảng cách giữa hai phân phối) buộc dự đoán của LSTM trong mạng sinh "
    "$p(z_t \\mid z_{1:t-1})$ – vốn không nhìn thấy dữ liệu – phải giống kết quả của mạng suy diễn. Nhờ phần này, mạng sinh "
    "học được quy luật chuyển cụm và có thể **tự dự báo** khi không còn dữ liệu (tương lai).",
])
p("Toàn bộ tham số (lớp tiền nội suy, hai LSTM, các MLP, tâm cụm $\\boldsymbol{\\mu}$) được học đồng thời bằng thuật toán "
  "Adam. Trong thực hành, hệ số của phần (2) được tăng dần từ 0 ở các epoch đầu (**KL annealing**) để mạng suy diễn kịp "
  "học phân cụm trước.")

h("3.7. Cơ chế cổng", 2)
p("Thay vì chọn tay $\\gamma$ trong ([[eq:psi]]), DGM² cho một mạng nhỏ tự quyết định $\\gamma$ ở từng bước dựa trên trạng "
  "thái của mạng suy diễn:")
eq(r"\gamma_t = \mathrm{sigmoid}\left(\mathrm{MLP}(\tilde{\mathbf{h}}_t)\right) \in (0, 1)", label="gate")
p("Khi mô hình tự tin về động học (chuỗi đang ổn định trong một cụm), $\\gamma_t$ nhỏ và dự báo dựa vào LSTM; khi không "
  "chắc chắn, $\\gamma_t$ lớn hơn và dự báo “dựa” vào hỗn hợp cơ sở chung.")

h("3.8. Dự báo", 2)
p("Quy trình dự báo $r$ bước cho một chuỗi mới:", indent=False)
R.numbered([
    "Tiền nội suy các ô thiếu của $w$ bước quá khứ (công thức ([[eq:impute]])).",
    "Mạng suy diễn đọc quá khứ, xác định chuỗi đang ở cụm nào; LSTM của mạng sinh đọc chuỗi cụm đó để có trạng thái $\\mathbf{h}_w$.",
    "Lặp cho $t = w, \\dots, w + r - 1$: tính $p(z_{t+1} \\mid z_{1:t})$ theo ([[eq:trans]]), cổng $\\gamma$ theo ([[eq:gate]]), "
    "trọng số $\\boldsymbol{\\psi}_{t+1}$ theo ([[eq:psi]]), rồi lấy **kỳ vọng** của hỗn hợp Gaussian làm giá trị dự báo:",
])
eq(r"\tilde{\mathbf{x}}_{t+1} = \sum_{c=1}^{k} \psi_{t+1}[c] \; \boldsymbol{\mu}_c", label="fc")
p("Dự báo là tổ hợp các tâm cụm học được từ toàn bộ dữ liệu, nên luôn nằm trong vùng giá trị hợp lý và ít bị ảnh "
  "hưởng khi quá khứ bị thiếu nhiều. Ngoài giá trị dự báo, phân phối $\\boldsymbol{\\psi}_{t+1}$ còn cho biết **độ bất định**: "
  "nếu xác suất dồn vào một cụm thì dự báo chắc chắn, nếu trải đều nhiều cụm thì kém chắc chắn.")

# =============================================================== CHƯƠNG 4
h("Chương 4. Demo: dự báo khí hậu trên dữ liệu thưa USHCN")
h("4.1. Dữ liệu và bài toán", 2)
p("Dữ liệu lấy từ mã nguồn của bài báo [2]: các chuỗi 100 ngày ghi "
  "5 biến khí hậu (lượng mưa, tuyết rơi, độ dày tuyết, nhiệt độ cao nhất, thấp nhất) của các trạm khí tượng Hoa Kỳ "
  "(USHCN), gồm 4000 chuỗi huấn luyện và 1000 chuỗi kiểm thử; khoảng **10,6%** giá trị bị thiếu. Nhiệm vụ: dùng "
  "**80 ngày** quá khứ để dự báo **20 ngày** tiếp theo. Dữ liệu được chuẩn hoá z-score theo từng biến, sai số tính "
  "trên thang chuẩn hoá (RMSE, MAE – càng nhỏ càng tốt).")
h("4.2. Cài đặt", 2)
p(f"DGM²-L được cài đặt lại bằng PyTorch trong khoảng 200 dòng (dgm2.py), đúng theo các công thức "
  f"([[eq:impute]])–([[eq:fc]]). Cấu hình: $k = {cfg['k']}$ cụm, LSTM kích thước {cfg['hidden']}, "
  f"phương sai $\\sigma^{{-1}} = {cfg['var']}$, dùng hàm cổng, huấn luyện {cfg['epochs']} epoch bằng Adam trên CPU "
  f"(khoảng 10 phút). Số cụm được chọn trên tập validation.")
table(["Tệp", "Chức năng"], [
    ["data.py", "Nạp và chuẩn hoá dữ liệu USHCN, tạo mặt nạ, xoá thêm quan sát"],
    ["dgm2.py", "Mô hình DGM²: tiền nội suy, mạng sinh, hỗn hợp động, mạng suy diễn, cổng, ELBO, dự báo"],
    ["train.py / evaluate.py", "Huấn luyện / đánh giá trên tập kiểm thử"],
    ["app.py", "Ứng dụng web tương tác (Streamlit)"],
], "Mã nguồn demo.", col_widths=[4.5, 11.5], font_size=11, label="code")

h("4.3. Kết quả dự báo", 2)
o = ev["overall"]; hz = ev["per_horizon_rmse"]; rb = ev["robustness"]
p(f"Trên 1000 chuỗi kiểm thử, DGM²-L đạt **RMSE = {f4(o['RMSE'])}** và **MAE = {f4(o['MAE'])}**. Sai số khá đều giữa "
  f"các biến (RMSE từ {f4(min(v['RMSE'] for v in ev['per_variable'].values()))} đến "
  f"{f4(max(v['RMSE'] for v in ev['per_variable'].values()))}). Hình [[fig:eval]](a) cho thấy sai số tăng dần theo tầm "
  f"dự báo, từ khoảng {f4(hz[0])} ở ngày đầu lên khoảng {f4(hz[-1])} ở ngày thứ 20 – càng xa hiện tại, mô hình càng phải "
  f"dựa vào quy luật chuyển cụm do LSTM học được.")
fig(os.path.join(RES, "forecast_example_123.png"),
    "Dự báo trên một chuỗi kiểm thử: chấm đen – quan sát thực; nét đứt cam – giá trị tiền nội suy; đường đỏ – dự báo "
    "(công thức ([[eq:fc]])); dải đỏ nhạt – độ bất định của hỗn hợp Gaussian; vùng xám – 20 ngày cần dự báo.", 15.5, label="fc")
fig(os.path.join(RES, "clusters_example_123.png"),
    "Cụm tiềm ẩn của cùng chuỗi: trên – xác suất thuộc từng cụm theo thời gian (trái vạch xanh: mạng suy diễn đọc quá khứ; "
    "phải vạch: trọng số $\\boldsymbol{\\psi}_t$ khi dự báo); dưới – giá trị cổng $\\gamma_t$.", 15.5, label="cl")
p("Hình [[fig:cl]] cho thấy cách mô hình “hiểu” chuỗi: khoảng 20 ngày đầu chuỗi dao động giữa cụm 11 và 20, sau đó "
  "chuyển hẳn sang cụm 35 và ở lại đó – mô hình đã tóm tắt 5 biến khí hậu thành một chuỗi **kiểu thời tiết** rời rạc. "
  "Khi dự báo, $\\boldsymbol{\\psi}_t$ tiếp tục dồn xác suất vào cụm 35, nên đường dự báo ở Hình [[fig:fc]] bám quanh tâm của "
  "cụm này. Hàm cổng học được $\\gamma_t \\approx 0$, đúng với nhận xét của bài báo rằng chỉ cần một lượng rất nhỏ hỗn hợp "
  "cơ sở.")

h("4.4. Độ bền khi dữ liệu thưa hơn", 2)
p("Để kiểm tra ưu điểm chính của DGM², ta giữ nguyên mô hình đã huấn luyện và xoá ngẫu nhiên thêm một phần dữ liệu "
  "quan sát được ở 80 ngày đầu vào của tập kiểm thử:")
table(["Tỉ lệ thiếu của đầu vào", "RMSE", "MAE"],
      [[f"{v['missing_ratio'] * 100:.0f}%", f4(v["RMSE"]), f4(v["MAE"])] for v in rb.values()],
      "Sai số dự báo khi dữ liệu đầu vào thưa dần.", col_widths=[5.5, 3.0, 3.0], label="rob")
fig(os.path.join(RES, "eval_horizon_robustness.png"),
    "(a) RMSE theo tầm dự báo 1–20 ngày; (b) sai số khi dữ liệu đầu vào thưa dần.", 15.5, label="eval")
r0, r60, r80 = rb["0.0"], rb["0.6"], rb["0.8"]
p(f"Khi tỉ lệ thiếu tăng từ {r0['missing_ratio'] * 100:.0f}% lên {r60['missing_ratio'] * 100:.0f}%, RMSE chỉ tăng từ "
  f"{f4(r0['RMSE'])} lên {f4(r60['RMSE'])} (khoảng {100 * (r60['RMSE'] / r0['RMSE'] - 1):.0f}%); kể cả khi hơn "
  f"{r80['missing_ratio'] * 100:.0f}% dữ liệu bị mất, RMSE vẫn là {f4(r80['RMSE'])}. Lý do: mô hình chỉ cần vài quan sát "
  f"để nhận ra chuỗi đang ở **cụm** nào, còn giá trị dự báo được sinh từ các tâm cụm học từ toàn bộ dữ liệu – đúng như "
  f"ý tưởng ở Chương 1.")

h("4.5. Ứng dụng web tương tác", 2)
p("Ứng dụng Streamlit (lệnh: cd demo, rồi streamlit run app.py) cho phép chọn một chuỗi kiểm thử, kéo thanh trượt để "
  "xoá bớt dữ liệu quá khứ, và xem ngay: dự báo 20 ngày kèm độ bất định, xác suất các cụm theo thời gian, giá trị cổng "
  "$\\gamma_t$ và vị trí các tâm cụm. Ứng dụng chỉ dùng mô hình đã huấn luyện sẵn nên chạy nhẹ, không cần GPU.")

# =============================================================== CHƯƠNG 5
h("Chương 5. Kết luận")
p("DGM² mô hình hoá chuỗi thời gian nhiều chiều thưa bằng cách kết hợp hai thành phần: **GMM** mô tả các cụm (trạng thái) "
  "chung của dữ liệu, và **LSTM** học quy luật chuyển giữa các cụm theo thời gian. Điểm then chốt là **hỗn hợp Gaussian "
  "động** – trọng số của GMM thay đổi theo từng thời điểm, pha trộn giữa dự đoán của LSTM và phân phối cơ sở chung, được "
  "điều tiết bởi một hàm cổng. Mô hình được huấn luyện bằng ELBO với một mạng suy diễn đóng vai trò gán cụm cho dữ liệu.")
p(f"Demo trên dữ liệu khí hậu USHCN cho kết quả dự báo 20 ngày với RMSE {f4(o['RMSE'])}, và sai số gần như không đổi khi "
  f"dữ liệu đầu vào thưa tới khoảng {r60['missing_ratio'] * 100:.0f}% – thể hiện đúng ưu điểm “bền vững với dữ liệu thiếu” "
  f"của mô hình. Ngoài dự báo, mô hình còn cho kết quả **dễ diễn giải**: chuỗi các kiểu thời tiết và độ bất định của dự "
  f"báo. Hạn chế: dự báo là tổ hợp của một số hữu hạn tâm cụm nên bị giới hạn độ chi tiết; hướng phát triển là dùng "
  f"phương sai riêng cho từng cụm hoặc tăng số cụm.")

# =============================================================== TÀI LIỆU
h("Tài liệu tham khảo")
refs = [
    "Wu, Y. et al. (2021). Dynamic Gaussian Mixture based Deep Generative Model for Robust Forecasting on Sparse "
    "Multivariate Time Series. AAAI 35(1): 651–659. https://ojs.aaai.org/index.php/AAAI/article/view/16145",
    "Mã nguồn và dữ liệu DGM²: https://github.com/KnowledgeDiscovery/DynamicGaussianMixture",
    "Bishop, C. M. (2006). Pattern Recognition and Machine Learning, Chương 9: Mixture Models and EM. Springer.",
    "Hochreiter, S.; Schmidhuber, J. (1997). Long Short-Term Memory. Neural Computation 9(8).",
    "Jang, E.; Gu, S.; Poole, B. (2017). Categorical Reparameterization with Gumbel-Softmax. ICLR.",
]
for i, r_ in enumerate(refs, 1):
    par = R.doc.add_paragraph(); par.paragraph_format.left_indent = Cm(0.8); par.paragraph_format.first_line_indent = Cm(-0.8)
    par.add_run(f"[{i}] {r_}").font.size = Pt(12)

out = os.path.join(HERE, "BaoCao_DGM2.docx")
try:
    R.save(out)
except PermissionError:          # file đang mở trong Word
    out = os.path.join(HERE, "BaoCao_DGM2_moi.docx")
    R.save(out)
print("saved", out, "| equations:", R.eq_no, "figures:", R.fig_no, "tables:", R.tab_no)
