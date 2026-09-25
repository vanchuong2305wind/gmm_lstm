"""Các hàm tiện ích dựng báo cáo Word bằng python-docx (công thức native OMML)."""
import copy, re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from omml import latex_to_omml

FONT = "Times New Roman"


class Report:
    def __init__(self):
        self.doc = Document()
        self.eq_no = 0
        self.fig_no = 0
        self.tab_no = 0
        self.labels = {}
        self._setup()

    # ------------------------------------------------------------ setup
    def _setup(self):
        d = self.doc
        sec = d.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.0)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        st = d.styles["Normal"]
        st.font.name, st.font.size = FONT, Pt(13)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.space_after, pf.line_spacing = Pt(6), 1.3
        for lvl, size in [(1, 16), (2, 14), (3, 13)]:
            h = d.styles[f"Heading {lvl}"]
            h.font.name, h.font.size, h.font.bold = FONT, Pt(size), True
            h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x68)
            rpr = h.element.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rfonts.set(qn(a), FONT)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rfonts.get(qn(a)) is not None:
                    del rfonts.attrib[qn(a)]
            h.italic = False
            h.paragraph_format.space_before = Pt(12 if lvl > 1 else 18)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        cap = d.styles["Caption"]
        cap.font.name, cap.font.size, cap.font.italic, cap.font.bold = FONT, Pt(11.5), True, False
        cap.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        self._page_numbers(sec)

    def _page_numbers(self, sec):
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._field(p, "PAGE")

    def _field(self, paragraph, instr):
        run = paragraph.add_run()
        for tag, text in [("begin", None), (None, instr), ("separate", None), (None, None), ("end", None)]:
            if tag:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); run._r.append(el)
            elif text:
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = f" {text} "; run._r.append(el)
        return run

    # ------------------------------------------------------------ text
    @staticmethod
    def _tokenize(text):
        """Tách văn bản thành các đoạn: $math$, **đậm**, *nghiêng*, thường (bỏ qua * nằm trong $...$)."""
        def find_close(s, start, marker):
            i = start
            while i < len(s):
                if s[i] == "$":
                    i = s.index("$", i + 1) + 1
                    continue
                if s.startswith(marker, i) and not (marker == "*" and s.startswith("**", i)):
                    return i
                i += 1
            return -1

        out, buf, i = [], "", 0
        while i < len(text):
            c = text[i]
            if c == "$":
                j = text.index("$", i + 1)
                out += [buf, text[i:j + 1]]; buf = ""; i = j + 1
            elif text.startswith("**", i) and find_close(text, i + 2, "**") > 0:
                j = find_close(text, i + 2, "**")
                out += [buf, text[i:j + 2]]; buf = ""; i = j + 2
            elif c == "*" and find_close(text, i + 1, "*") > 0:
                j = find_close(text, i + 1, "*")
                out += [buf, text[i:j + 1]]; buf = ""; i = j + 1
            else:
                buf += c; i += 1
        out.append(buf)
        return [t for t in out if t]

    def _add_rich(self, p, text, size=None, bold=False, italic=False):
        """Văn bản có **đậm**, *nghiêng* và $công thức nội dòng$ (công thức được phép nằm trong đậm/nghiêng)."""
        for tok in self._tokenize(text):
            if not tok:
                continue
            if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
                self._add_rich(p, tok[2:-2], size, True, italic)
            elif tok.startswith("$") and tok.endswith("$"):
                p._p.append(copy.deepcopy(latex_to_omml(tok[1:-1])))
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                self._add_rich(p, tok[1:-1], size, bold, True)
            else:
                r = p.add_run(tok); r.bold = bold or None; r.italic = italic or None
                if size: r.font.size = Pt(size)
        return p

    def h(self, text, level=1):
        return self.doc.add_heading(text, level)

    def p(self, text, align="justify", indent=True, size=None, space_after=None):
        par = self.doc.add_paragraph()
        par.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                         "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        if indent and align == "justify":
            par.paragraph_format.first_line_indent = Cm(1.0)
        if space_after is not None:
            par.paragraph_format.space_after = Pt(space_after)
        return self._add_rich(par, text, size)

    def bullets(self, items, style="List Bullet"):
        for it in items:
            par = self.doc.add_paragraph(style=style)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.space_after = Pt(3)
            self._add_rich(par, it)

    def numbered(self, items):
        self.bullets(items, style="List Number")

    def eq(self, latex, number=True, label=None):
        """Công thức hiển thị (display): bảng 2 cột không viền – cột trái công thức căn giữa, cột phải số (n)."""
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        width = self.doc.sections[0].page_width - self.doc.sections[0].left_margin - self.doc.sections[0].right_margin
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        wnum = Cm(1.4)
        for cell, w in zip(t.rows[0].cells, [width - wnum, wnum]):
            cell.width = w
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        c0, c1 = t.rows[0].cells
        par = c0.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(2); par.paragraph_format.space_after = Pt(2)
        omp = OxmlElement("m:oMathPara")
        omp.append(copy.deepcopy(latex_to_omml(latex)))
        par._p.append(omp)
        pn = c1.paragraphs[0]; pn.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pn.paragraph_format.space_after = Pt(0)
        if number:
            self.eq_no += 1
            if label: self.labels['eq:' + label] = self.eq_no
            pn.add_run(f"({self.eq_no})")
        return self.eq_no

    # ------------------------------------------------------------ figures / tables
    def fig(self, path, caption, width_cm=15.5, label=None):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(path, width=Cm(width_cm))
        self.fig_no += 1
        if label: self.labels['fig:' + label] = self.fig_no
        c = self.doc.add_paragraph(style="Caption")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(f"Hình {self.fig_no}. "); r.bold = True
        self._add_rich(c, caption)
        return self.fig_no

    def table(self, header, rows, caption, col_widths=None, bold_best=None, font_size=11, label=None):
        self.tab_no += 1
        if label: self.labels['tab:' + label] = self.tab_no
        c = self.doc.add_paragraph(style="Caption")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.keep_with_next = True
        r = c.add_run(f"Bảng {self.tab_no}. "); r.bold = True
        self._add_rich(c, caption)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, hd in enumerate(header):
            cell = t.rows[0].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]; par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._add_rich(par, f"**{hd}**" if "$" not in hd else hd, size=font_size)
            self._shade(cell, "DCE6F1")
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.rows[i + 1].cells[j]
                cell.text = ""
                par = cell.paragraphs[0]
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
                par.paragraph_format.space_after = Pt(0)
                txt = str(val)
                if bold_best and (i, j) in bold_best:
                    txt = f"**{txt}**"
                self._add_rich(par, txt, size=font_size)
        if col_widths:
            for row in t.rows:
                for j, w in enumerate(col_widths):
                    row.cells[j].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return self.tab_no

    @staticmethod
    def _shade(cell, color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
        tcPr.append(shd)

    def code(self, text, size=9.5):
        for line in text.strip("\n").split("\n"):
            par = self.doc.add_paragraph()
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.left_indent = Cm(0.5)
            r = par.add_run(line if line else " ")
            r.font.name, r.font.size = "Consolas", Pt(size)
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Consolas")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def toc(self):
        par = self.doc.add_paragraph()
        self._field(par, 'TOC \\o "1-3" \\h \\z \\u')

    def resolve_refs(self):
        """Thay các chỗ giữ [[eq:x]], [[fig:x]], [[tab:x]] bằng số thứ tự thật."""
        pat = re.compile(r"\[\[((?:eq|fig|tab):[\w-]+)\]\]")
        for t in self.doc.element.body.iter(qn("w:t")):
            if t.text and "[[" in t.text:
                def sub(mo):
                    assert mo.group(1) in self.labels, "Thiếu nhãn " + mo.group(1)
                    return str(self.labels[mo.group(1)])
                t.text = pat.sub(sub, t.text)

    def save(self, path):
        self.resolve_refs()
        # đặt chế độ tương thích Word 2013+ để Word không mở ở "Compatibility Mode"
        settings = self.doc.settings.element
        compat = settings.find(qn("w:compat"))
        if compat is None:
            compat = OxmlElement("w:compat"); settings.append(compat)
        for cs in compat.findall(qn("w:compatSetting")):
            if cs.get(qn("w:name")) == "compatibilityMode":
                compat.remove(cs)
        cs = OxmlElement("w:compatSetting")
        for k, v in [("w:name", "compatibilityMode"), ("w:uri", "http://schemas.microsoft.com/office/word"), ("w:val", "15")]:
            cs.set(qn(k), v)
        compat.append(cs)
        self.doc.save(path)
